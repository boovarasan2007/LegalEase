import os
import re
from typing import Dict, Any

def normalize_text(text: str) -> str:
    """Cleans extra whitespace, normalizes line endings, and preserves paragraph breaks."""
    if not text:
        return ""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_pdf_text(filepath: str) -> Dict[str, Any]:
    """Extracts text page by page from PDF using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        page_texts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_texts.append(f"--- Page {i + 1} ---\n" + page_text)
        full_text = "\n\n".join(page_texts)
        cleaned = normalize_text(full_text)
        page_count = len(reader.pages)
        return {
            "text": cleaned,
            "page_count": page_count,
            "word_count": len(cleaned.split()) if cleaned else 0,
            "char_count": len(cleaned),
            "is_empty": len(cleaned.replace("---", "").strip()) < 30
        }
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_docx_text(filepath: str) -> Dict[str, Any]:
    """Extracts paragraphs and tables from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(filepath)
        parts = []

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        full_text = "\n\n".join(parts)
        cleaned = normalize_text(full_text)
        word_count = len(cleaned.split()) if cleaned else 0
        est_pages = max(1, (word_count // 400) + 1)
        return {
            "text": cleaned,
            "page_count": est_pages,
            "word_count": word_count,
            "char_count": len(cleaned),
            "is_empty": len(cleaned.strip()) < 20
        }
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def extract_txt_text(filepath: str) -> Dict[str, Any]:
    """Reads text file with UTF-8 and latin-1 fallback."""
    try:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                raw_text = f.read()
        except UnicodeDecodeError:
            with open(filepath, 'r', encoding='latin-1') as f:
                raw_text = f.read()

        cleaned = normalize_text(raw_text)
        word_count = len(cleaned.split()) if cleaned else 0
        est_pages = max(1, (word_count // 400) + 1)
        return {
            "text": cleaned,
            "page_count": est_pages,
            "word_count": word_count,
            "char_count": len(cleaned),
            "is_empty": len(cleaned.strip()) < 10
        }
    except Exception as e:
        raise ValueError(f"Failed to read TXT file: {str(e)}")

def extract_document_text(filepath: str, original_filename: str) -> Dict[str, Any]:
    """High-level dispatcher for PDF, DOCX, and TXT files."""
    ext = os.path.splitext(original_filename)[1].lower().replace('.', '')

    if ext == 'pdf':
        res = extract_pdf_text(filepath)
    elif ext == 'docx':
        res = extract_docx_text(filepath)
    elif ext in ['txt', 'text']:
        res = extract_txt_text(filepath)
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Supported formats: PDF, DOCX, TXT.")

    if res.get("is_empty"):
        raise ValueError(
            "LegalEase could not extract readable text from this document. "
            "The file may contain scanned images. OCR support can be added in a future version."
        )

    return res
